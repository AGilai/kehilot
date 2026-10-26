"""Bounded, atomic, additive Excel imports of the association phone list."""
import base64
from io import BytesIO
import zipfile

from openpyxl import load_workbook, Workbook

from .phone_auth import mobile, assign_membership
from .validation import APIError, text, israel_today

MAX_BYTES = 5 * 1024 * 1024
HEADERS = {
    'שם פרטי': 'first', 'firstName': 'first', 'first_name': 'first',
    'שם משפחה': 'last', 'lastName': 'last', 'last_name': 'last',
    'שם': 'name', 'שם מלא': 'name', 'name': 'name',
    'פלאפון': 'phone', 'טלפון': 'phone', 'מספר פלאפון': 'phone', 'מספר טלפון': 'phone', 'phone': 'phone',
}


def parse_workbook(content):
    if not isinstance(content, bytes) or not 0 < len(content) <= MAX_BYTES:
        raise APIError(413, 'גודל קובץ האקסל המרבי הוא 5MB.')
    workbook = None
    try:
        with zipfile.ZipFile(BytesIO(content)) as archive:
            entries = archive.infolist()
            if len(entries) > 200 or sum(item.file_size for item in entries) > 20 * 1024 * 1024:
                raise APIError(400, 'קובץ האקסל גדול מדי לאחר פתיחה.')
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=False, keep_links=False)
        sheet = workbook.worksheets[0]
        # Do not trust worksheet dimensions from a third-party file.
        sheet.reset_dimensions()
        rows = sheet.iter_rows(max_row=10002, max_col=20)
        header = next(rows)
        mapping = {}
        for index, cell in enumerate(header):
            field = HEADERS.get(str(cell.value or '').strip())
            if field:
                if field in mapping:
                    raise APIError(400, 'בקובץ יש כותרת עמודה כפולה.')
                mapping[field] = index
        if 'phone' not in mapping or not ('name' in mapping or {'first', 'last'} <= mapping.keys()):
            raise APIError(400, 'נדרשות כותרות: שם פרטי, שם משפחה, פלאפון; או שם מלא ופלאפון.')
        result, seen = [], set()
        for number, cells in enumerate(rows, 2):
            if number > 10001:
                raise APIError(400, 'ניתן לייבא עד 10,000 שורות.')
            if not any(cell.value is not None for cell in cells):
                continue
            if any(cell.data_type == 'f' for cell in cells):
                raise APIError(400, f'שורה {number}: יש להזין ערכים ולא נוסחאות.')
            def value(field):
                raw = cells[mapping[field]].value if field in mapping else None
                if isinstance(raw, bool):
                    raise APIError(400, f'שורה {number}: ערך אינו תקין.')
                return str(int(raw)) if isinstance(raw, (int, float)) and float(raw).is_integer() else str(raw or '').strip()
            try:
                raw_phone = value('phone')
                # Excel often drops the initial zero from a numeric Israeli phone.
                if raw_phone.isdigit() and len(raw_phone) == 9 and raw_phone.startswith('5'):
                    raw_phone = '0' + raw_phone
                phone = mobile(raw_phone)
                first, last = text(value('first'), 'שם פרטי'), text(value('last'), 'שם משפחה')
                name = text((first + ' ' + last).strip() or value('name'), 'שם מלא', required=True)
            except APIError as error:
                raise APIError(400, f'שורה {number}: {error}') from None
            if phone in seen:
                raise APIError(400, f'שורה {number}: מספר פלאפון כפול בקובץ.')
            seen.add(phone)
            result.append(dict(phone=phone, first=first, last=last, name=name))
        if not result:
            raise APIError(400, 'לא נמצאו חברי עמותה בקובץ.')
        return result
    except APIError:
        raise
    except Exception:
        raise APIError(400, 'לא ניתן לקרוא את קובץ האקסל. יש לבחור קובץ ‎.xlsx תקין.') from None
    finally:
        if workbook:
            workbook.close()


def template(path):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = 'חברי עמותה'
    sheet.sheet_view.rightToLeft = True
    sheet.append(['שם פרטי', 'שם משפחה', 'פלאפון'])
    for column in ('A', 'B', 'C'):
        sheet.column_dimensions[column].width = 24
    # Phone column is text so Excel preserves the leading zero.
    for row in range(2, 1002):
        sheet.cell(row, 3).number_format = '@'
    sheet.freeze_panes = 'A2'
    workbook.save(path)
    workbook.close()


class MemberImportStore:
    def association_list(self, user):
        self.require_gabbai(user)
        with self.db() as db:
            return [dict(phone=row['phone'], name=row['display_name'], firstName=row['first_name'],
                         lastName=row['last_name'], registered=bool(row['user_id']))
                    for row in db.execute('''SELECT a.*,v.user_id FROM association_members a
                        LEFT JOIN verified_phones v ON v.phone=a.phone ORDER BY a.last_name,a.first_name,a.phone''')]

    def import_members(self, user, value):
        from .store import now
        self.require_gabbai(user)
        if not isinstance(value, dict) or set(value) != {'name', 'base64'} or not isinstance(value['name'], str) or not value['name'].lower().endswith('.xlsx'):
            raise APIError(400, 'יש לבחור קובץ ‎.xlsx.')
        encoded = value['base64']
        if not isinstance(encoded, str) or len(encoded) > ((MAX_BYTES + 2) // 3) * 4:
            raise APIError(413, 'גודל קובץ האקסל המרבי הוא 5MB.')
        try:
            content = base64.b64decode(encoded, validate=True)
        except ValueError:
            raise APIError(400, 'קובץ האקסל אינו תקין.') from None
        entries = parse_workbook(content)
        added, updated = 0, 0
        with self.db() as db:
            db.execute('BEGIN IMMEDIATE')
            for entry in entries:
                old = db.execute('SELECT 1 FROM association_members WHERE phone=?', (entry['phone'],)).fetchone()
                added += int(not old)
                updated += int(bool(old))
                db.execute('''INSERT INTO association_members VALUES (?,?,?,?,?,?)
                    ON CONFLICT(phone) DO UPDATE SET first_name=excluded.first_name,last_name=excluded.last_name,
                    display_name=excluded.display_name,imported_at=excluded.imported_at,imported_by=excluded.imported_by''',
                    (entry['phone'], entry['first'], entry['last'], entry['name'], now(), user['id']))
                linked = db.execute('SELECT user_id FROM verified_phones WHERE phone=?', (entry['phone'],)).fetchone()
                if linked:
                    assign_membership(db, linked['user_id'])
            self._refresh_roster(db, israel_today())
        return dict(added=added, updated=updated, total=len(entries))
