"""Permanent physical seat layout with per-prayer, current-period availability."""
from collections import Counter, defaultdict
from datetime import datetime, timezone
import json
import re

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from .seating import current_time, event_value
from .validation import APIError

FREE_COLOR = '#6AA84F'
SEAT_PATTERN = re.compile(r'(.+?)[\-־–—]\s*(\d+)\s*$')


def initialize(db):
    db.executescript('''
        CREATE TABLE IF NOT EXISTS seating_map_layout (
            id INTEGER PRIMARY KEY CHECK(id=1), revision INTEGER NOT NULL,
            layout TEXT NOT NULL, imported_at TEXT NOT NULL);
        CREATE TABLE IF NOT EXISTS seating_map_assignments (
            cell_id TEXT PRIMARY KEY,
            person_id TEXT REFERENCES roster_people(id),
            owner_label TEXT NOT NULL, serial INTEGER NOT NULL CHECK(serial>0),
            UNIQUE(person_id,serial));
    ''')


def color(value, fallback):
    return '#' + value.rgb[-6:] if value is not None and value.type == 'rgb' and isinstance(value.rgb, str) else fallback


def read_layout(path, sheet='Map', overrides=None):
    book = load_workbook(path, data_only=True, keep_links=False)
    try:
        if sheet not in book.sheetnames:
            raise APIError(400, 'גיליון המפה לא נמצא.')
        source = book[sheet]
        for coordinate, change in (overrides or {}).items():
            if source[coordinate].value != change['from']:
                raise APIError(409, 'תיקון משבצת אינו תואם לקובץ המקור: ' + coordinate)
            source[coordinate] = change['to']
        meaningful = [cell for row in source for cell in row if cell.value is not None and str(cell.value).strip()]
        if not meaningful:
            raise APIError(400, 'מפת המקומות ריקה.')
        min_row, max_row = min(c.row for c in meaningful), max(c.row for c in meaningful)
        min_col, max_col = min(c.column for c in meaningful), max(c.column for c in meaningful)
        merges = {}
        covered = set()
        for merged in source.merged_cells.ranges:
            if source.cell(merged.min_row, merged.min_col).value is None:
                continue
            min_row, max_row = min(min_row, merged.min_row), max(max_row, merged.max_row)
            min_col, max_col = min(min_col, merged.min_col), max(max_col, merged.max_col)
            merges[(merged.min_row, merged.min_col)] = (merged.max_row, merged.max_col)
            covered.update((r, c) for r in range(merged.min_row, merged.max_row + 1) for c in range(merged.min_col, merged.max_col + 1)
                           if (r, c) != (merged.min_row, merged.min_col))
        if (max_row - min_row + 1) * (max_col - min_col + 1) > 10000:
            raise APIError(400, 'שטח המפה גדול מדי. יש לצמצם את הגיליון לאזור המושבים.')
        widths = []
        for column in range(min_col, max_col + 1):
            dimension = next((d for d in source.column_dimensions.values() if d.min <= column <= d.max), None)
            characters = dimension.width if dimension else source.sheet_format.defaultColWidth or 8.43
            widths.append(0 if dimension and dimension.hidden else (characters * 7 + 5) * .75)
        heights = [0 if source.row_dimensions.get(row) and source.row_dimensions[row].hidden else
                   (source.row_dimensions[row].height if source.row_dimensions.get(row) and source.row_dimensions[row].height else source.sheet_format.defaultRowHeight or 15)
                   for row in range(min_row, max_row + 1)]
        xs, ys = [0], [0]
        for value in widths: xs.append(xs[-1] + value)
        for value in heights: ys.append(ys[-1] + value)
        cells = []
        for row in range(min_row, max_row + 1):
            for column in range(min_col, max_col + 1):
                if (row, column) in covered:
                    continue
                cell = source.cell(row, column)
                label = str(cell.value).strip() if cell.value is not None else ''
                if len(label) > 200:
                    raise APIError(400, 'כיתוב במפה ארוך מדי: ' + cell.coordinate)
                borders = [bool(getattr(cell.border, side) and getattr(cell.border, side).style) for side in ('top','right','bottom','left')]
                if not label and not any(borders) and cell.fill.patternType != 'solid':
                    continue
                last_row, last_col = merges.get((row, column), (row, column))
                x, y = xs[column - min_col], ys[row - min_row]
                width, height = xs[last_col - min_col + 1] - x, ys[last_row - min_row + 1] - y
                if label and (width <= 0 or height <= 0):
                    raise APIError(400, 'יש משבצת מוסתרת במפה: ' + cell.coordinate)
                match = SEAT_PATTERN.fullmatch(label)
                kind = 'seat' if match else 'free' if label == 'פנוי' else 'label'
                value = dict(id=cell.coordinate, x=round(xs[-1] - x - width if source.sheet_view.rightToLeft else x, 3),
                             y=round(y, 3), width=round(width, 3), height=round(height, 3), label=label, kind=kind,
                             background=color(cell.fill.fgColor, '#FFFFFF') if cell.fill.patternType == 'solid' else '#FFFFFF',
                             color=color(cell.font.color, '#182A23'), fontSize=cell.font.sz or 11, bold=bool(cell.font.b), borders=borders)
                if match:
                    value.update(ownerLabel=match[1].strip(), serial=int(match[2]))
                cells.append(value)
        for owner in {c['ownerLabel'] for c in cells if c['kind'] == 'seat'}:
            numbers = sorted(c['serial'] for c in cells if c.get('ownerLabel') == owner)
            if numbers != list(range(1, len(numbers) + 1)):
                raise APIError(400, 'מספור המקומות חסר או כפול עבור ' + owner)
        return dict(width=round(xs[-1], 3), height=round(ys[-1], 3), cells=cells,
                    sourceRange=f'{get_column_letter(min_col)}{min_row}:{get_column_letter(max_col)}{max_row}',
                    sheet=sheet, rightToLeft=bool(source.sheet_view.rightToLeft))
    finally:
        book.close()


def import_layout(db, layout, owners, changes=(), apply=False):
    labels = {cell['ownerLabel'] for cell in layout['cells'] if cell['kind'] == 'seat'}
    if set(owners) != labels:
        raise APIError(400, 'יש לאשר שיוך לכל בעלי המקומות במפה, ללא שמות נוספים.')
    promotions = {}
    for change in changes:
        if set(change) != {'personId','from','to'} or change['to'] != 'א' or change['personId'] in promotions:
            raise APIError(400, 'שינוי סיווג בייבוא אינו תקין.')
        row = db.execute('SELECT classification FROM roster_people WHERE id=? AND active=1', (change['personId'],)).fetchone()
        if not row or row['classification'] not in (change['from'], change['to']):
            raise APIError(409, 'סיווג המתפלל אינו תואם לשינוי שאושר.')
        promotions[change['personId']] = change
    assignments, totals, serials = [], Counter(), set()
    for cell in layout['cells']:
        if cell['kind'] != 'seat':
            continue
        person_id = owners[cell['ownerLabel']]
        if person_id is not None:
            person = db.execute('SELECT * FROM roster_people WHERE id=? AND active=1', (person_id,)).fetchone()
            if not person or (person['classification'] != 'א' and person_id not in promotions):
                raise APIError(400, 'בעל מקום חייב להיות ראש משפחה פעיל בסיווג א: ' + cell['ownerLabel'])
            if (person_id, cell['serial']) in serials:
                raise APIError(400, 'מספר מקום כפול לאותו ראש משפחה.')
            serials.add((person_id, cell['serial'])); totals[person_id] += 1
        assignments.append((cell['id'], person_id, cell['ownerLabel'], cell['serial']))
    if set(promotions) - set(totals):
        raise APIError(400, 'שינוי סיווג ללא מקומות משויכים אינו מותר בייבוא מפה.')
    if apply:
        stamp = datetime.now(timezone.utc).isoformat()
        for identifier, change in promotions.items():
            db.execute("UPDATE roster_people SET classification='א',classification_override='א',updated_at=? WHERE id=?", (stamp, identifier))
        for identifier, count in totals.items():
            db.execute('''UPDATE roster_people SET allocated_seats=?,seat_allocation_revision=seat_allocation_revision+1,updated_at=?
                WHERE id=? AND (allocated_seats IS NULL OR allocated_seats!=?)''', (count, stamp, identifier, count))
        db.execute('DELETE FROM seating_map_assignments')
        db.executemany('INSERT INTO seating_map_assignments VALUES (?,?,?,?)', assignments)
        db.execute('''INSERT INTO seating_map_layout VALUES (1,1,?,?) ON CONFLICT(id) DO UPDATE SET
            revision=seating_map_layout.revision+1,layout=excluded.layout,imported_at=excluded.imported_at''',
            (json.dumps(layout, ensure_ascii=False), stamp))
    return dict(assignedSeats=len(assignments), freeSeats=sum(c['kind']=='free' for c in layout['cells']),
                households=len(totals), fixedSeats=sum(row[1] is None for row in assignments), classificationChanges=len(promotions))


class SeatingMapStore:
    def seating_map(self, user):
        self.require_gabbai(user)
        now = current_time()
        with self.db() as db:
            db.execute('BEGIN IMMEDIATE')
            event = self._current_seating_event(db, now)
            saved = db.execute('SELECT * FROM seating_map_layout WHERE id=1').fetchone()
            if not saved:
                return dict(layout=None, event=event_value(event, now), maps=[], warnings=[])
            layout = json.loads(saved['layout']); layout['revision'] = saved['revision']
            assignments = {row['cell_id']: row for row in db.execute('SELECT * FROM seating_map_assignments')}
            grouped = defaultdict(list)
            for row in assignments.values():
                if row['person_id']: grouped[row['person_id']].append(row)
            warnings, released = [], {prayer: set() for prayer in ('arvit','shacharit')}
            for identifier, seats in grouped.items():
                person = db.execute('SELECT * FROM roster_people WHERE id=?', (identifier,)).fetchone()
                status = self._seating_status(db, event, person, now)
                if not status['eligible'] or status['allocatedSeats'] != len(seats):
                    warnings.append('יש לבדוק את שיוך המקומות והמכסה של ' + seats[0]['owner_label'])
                    continue
                descending = sorted(seats, key=lambda row: row['serial'], reverse=True)
                for prayer in released:
                    released[prayer].update(row['cell_id'] for row in descending[:status['releasedSeats'][prayer]])
            for owner in sorted({row['owner_label'] for row in assignments.values() if row['person_id'] is None}):
                warnings.append('המקומות של ' + owner + ' שמורים ללא שיוך לראש משפחה.')
            maps = []
            for prayer, title in [('arvit','ערבית'),('shacharit','שחרית')]:
                cells = [dict(cellId=cell['id'], available=cell['kind']=='free' or cell['id'] in released[prayer],
                              label='פנוי' if cell['kind']=='free' or cell['id'] in released[prayer] else cell['label']) for cell in layout['cells']]
                maps.append(dict(prayer=prayer, title=title, cells=cells, freeSeats=sum(c['available'] for c in cells)))
            return dict(layout=layout, event=event_value(event, now), maps=maps, warnings=warnings,
                        generatedAt=now.isoformat(), freeColor=FREE_COLOR)
