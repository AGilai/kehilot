from datetime import datetime, timezone
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

from openpyxl import Workbook
from openpyxl.styles import PatternFill, Border, Side
import test_seating as fixtures
from server.seating_map import read_layout, import_layout
from server.validation import APIError


class SeatingMapTest(unittest.TestCase):
    setUpClass = classmethod(fixtures.SeatingTest.setUpClass.__func__)
    tearDownClass = classmethod(fixtures.SeatingTest.tearDownClass.__func__)
    setUp = fixtures.SeatingTest.setUp
    tearDown = fixtures.SeatingTest.tearDown
    http = fixtures.SeatingTest.http
    allocate = fixtures.SeatingTest.allocate
    status = fixtures.SeatingTest.status
    payload = staticmethod(fixtures.SeatingTest.payload)
    report = fixtures.SeatingTest.report
    special = fixtures.SeatingTest.special

    def workbook(self):
        book = Workbook(); sheet = book.active; sheet.title = 'Map'
        sheet.column_dimensions['A'].width = 10
        sheet.column_dimensions['B'].width = 8
        sheet.row_dimensions[2].height = 25
        for number, coordinate in [(1,'A2'),(2,'B2'),(3,'A3'),(4,'B3')]:
            sheet[coordinate] = 'משפחה-' + str(number)
            sheet[coordinate].fill = PatternFill('solid', fgColor='FFFF7C80')
            sheet[coordinate].border = Border(bottom=Side(style='thin', color='000000'))
        sheet['D2'] = 'פנוי'; sheet['D2'].fill = PatternFill('solid', fgColor='FF6AA84F')
        sheet.merge_cells('F2:G3'); sheet['F2'] = 'בימה'
        path = Path(self.temp.name)/'map.xlsx'; book.save(path); book.close()
        return path

    def install(self):
        layout = read_layout(self.workbook())
        with self.store.db() as db:
            import_layout(db, layout, {'משפחה':self.member_id}, apply=True)
        return layout

    def map(self):
        status, result, _ = self.http('GET', '/api/admin/seating-map', token='gabbai-token')
        self.assertEqual(status, 200, result)
        return result

    @staticmethod
    def free(report, prayer):
        return {cell['cellId'] for item in report['maps'] if item['prayer']==prayer for cell in item['cells'] if cell['available']}

    def test_excel_geometry_aisles_merge_and_colors_are_retained(self):
        layout = read_layout(self.workbook())
        cells = {cell['id']:cell for cell in layout['cells']}
        self.assertEqual(layout['sourceRange'], 'A2:G3')
        self.assertEqual(len(cells), 6)
        self.assertEqual(cells['A2']['height'], 25)
        self.assertGreater(cells['A2']['width'], cells['B2']['width'])
        self.assertGreater(cells['D2']['x'], cells['B2']['x'] + cells['B2']['width'])
        self.assertEqual(cells['F2']['height'], 40)
        self.assertEqual(cells['F2']['kind'], 'label')
        self.assertEqual(cells['A2']['background'], '#FF7C80')
        self.assertEqual(cells['D2']['background'], '#6AA84F')

    def test_highest_numbers_release_independently_for_each_prayer(self):
        self.install(); self.report(2, shacharit=1)
        result = self.map()
        self.assertEqual(self.free(result, 'arvit'), {'A3','B3','D2'})
        self.assertEqual(self.free(result, 'shacharit'), {'B3','D2'})
        labels = {c['cellId']:c['label'] for c in result['maps'][0]['cells']}
        self.assertEqual(labels['A2'], 'משפחה-1')
        self.assertEqual(labels['A3'], 'פנוי')
        self.assertEqual(labels['F2'], 'בימה')
        self.assertEqual(result['maps'][0]['freeSeats'], 3)
        self.assertEqual(result['warnings'], [])

    def test_ordinary_defaults_hold_seats_and_special_defaults_release_them(self):
        self.install()
        self.assertEqual(self.free(self.map(),'arvit'), {'D2'})
        self.special()
        self.assertEqual(self.free(self.map(),'arvit'), {'A2','B2','A3','B3','D2'})
        self.report(2, shacharit=3)
        self.assertEqual(self.free(self.map(),'arvit'), {'A3','B3','D2'})
        self.assertEqual(self.free(self.map(),'shacharit'), {'B3','D2'})

    def test_edits_and_period_rollover_do_not_mutate_the_permanent_map(self):
        original = self.install()
        self.report(3); self.report(1, shacharit=0)
        self.assertEqual(self.free(self.map(),'arvit'), {'B3','D2'})
        with patch('server.seating_map.current_time', return_value=datetime(2026,11,1,10,tzinfo=timezone.utc)):
            result = self.map()
        self.assertEqual(result['event']['date'], '2026-11-07')
        self.assertEqual(self.free(result,'arvit'), {'D2'})
        self.assertEqual(result['layout']['cells'], original['cells'])
        with self.store.db() as db:
            self.assertEqual(db.execute('SELECT COUNT(*) FROM seating_map_assignments').fetchone()[0], 4)
            self.assertEqual(db.execute('SELECT COUNT(*) FROM seating_reports').fetchone()[0], 0)

    def test_only_gabbai_can_read_map_and_unloaded_map_has_an_empty_state(self):
        self.assertEqual(self.http('GET','/api/admin/seating-map')[0], 401)
        self.assertEqual(self.http('GET','/api/admin/seating-map',token='member-token')[0], 403)
        self.assertIsNone(self.map()['layout'])

    def test_quota_mismatch_warns_and_never_releases_uncertain_seats(self):
        self.install(); self.report(2)
        self.allocate(self.member_id, 3)
        result = self.map()
        self.assertEqual(self.free(result,'arvit'), {'D2'})
        self.assertEqual(len(result['warnings']), 1)

    def test_import_requires_all_confirmed_owners_and_rejects_unapproved_classification(self):
        layout = read_layout(self.workbook())
        for mapping in [{}, {'משפחה':'missing'}, {'משפחה':self.other_id}]:
            with self.store.db() as db:
                with self.assertRaises(APIError): import_layout(db, layout, mapping, apply=True)
                self.assertEqual(db.execute('SELECT COUNT(*) FROM seating_map_assignments').fetchone()[0], 0)
        with self.store.db() as db:
            change = dict(personId=self.other_id, **{'from':'','to':'א'})
            import_layout(db, layout, {'משפחה':self.other_id}, [change], apply=True)
            person = db.execute('SELECT * FROM roster_people WHERE id=?',(self.other_id,)).fetchone()
            self.assertEqual(person['classification'], 'א')
            self.assertEqual(person['allocated_seats'], 4)

    def test_reimport_preserves_allocation_revision_if_counts_are_unchanged(self):
        layout = self.install(); self.report(1)
        before = self.status()['allocationRevision']
        with self.store.db() as db: import_layout(db, layout, {'משפחה':self.member_id}, apply=True)
        self.assertEqual(self.status()['allocationRevision'], before)
        self.assertTrue(self.status()['reported'])
        self.assertEqual(self.map()['layout']['revision'], 2)

    def test_overrides_are_checked_and_numbering_errors_are_not_silently_accepted(self):
        path = self.workbook()
        with self.assertRaises(APIError): read_layout(path, overrides={'A2':{'from':'wrong','to':'משפחה-5'}})
        with self.assertRaises(APIError): read_layout(path, overrides={'A2':{'from':'משפחה-1','to':'משפחה-4'}})
        layout = read_layout(path, overrides={'A2':{'from':'משפחה-1','to':'משפחה-2'},'B2':{'from':'משפחה-2','to':'משפחה-1'}})
        self.assertEqual(next(c['label'] for c in layout['cells'] if c['id']=='A2'),'משפחה-2')


if __name__ == '__main__': unittest.main()
